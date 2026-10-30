from datetime import date
from io import BytesIO

from django.contrib.auth.models import Group
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import Usuario
from core.permissions import ROL_ADMINISTRADOR, ROL_COORDINADOR, ROL_SECRETARIO
from members.models import Miembro


class ReportesYDashboardTests(TestCase):
    def setUp(self):
        Miembro.objects.create(nombre_completo="José Álvarez", fecha_ingreso=date.today())
        self.usuarios = {}
        for rol in (ROL_ADMINISTRADOR, ROL_SECRETARIO, ROL_COORDINADOR):
            usuario = Usuario.objects.create_user(f"{rol.lower()}@reportes.test", "secreto123", nombre=rol)
            usuario.groups.add(Group.objects.get(name=rol))
            self.usuarios[rol] = usuario

    def test_dashboard_renderiza_y_no_expone_finanzas(self):
        self.client.force_login(self.usuarios[ROL_ADMINISTRADOR])
        response = self.client.get(reverse("inicio"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Miembros activos")
        self.client.force_login(self.usuarios[ROL_SECRETARIO])
        response = self.client.get(reverse("inicio"))
        self.assertNotContains(response, "Ingresos del mes")
        self.assertNotContains(response, "ingresos-chart")

    def test_reporte_miembros_y_excel_con_filtros(self):
        self.client.force_login(self.usuarios[ROL_ADMINISTRADOR])
        response = self.client.get(reverse("reportes"), {"tipo": "miembros", "busqueda": "José"})
        self.assertContains(response, "José Álvarez")
        response = self.client.get(reverse("reportes"), {"tipo": "miembros", "busqueda": "José", "export": "excel"})
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        libro = load_workbook(filename=BytesIO(response.content))
        self.assertEqual(libro.active.cell(3, 1).value, "Nombre")

    def test_secretario_y_coordinador_no_acceden_a_ingresos(self):
        for rol in (ROL_SECRETARIO, ROL_COORDINADOR):
            self.client.force_login(self.usuarios[rol])
            self.assertEqual(self.client.get(reverse("reportes"), {"tipo": "ingresos"}).status_code, 403)

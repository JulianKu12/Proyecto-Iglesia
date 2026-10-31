from io import BytesIO
import os

_FONT = "Vera"


def excel_response(title, headers, rows, filename, iglesia):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    ws.append([iglesia])
    ws.append([title])
    ws.append(headers)
    for row in rows:
        ws.append([str(value) if value is not None else "" for value in row])
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4E89")
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{chr(64 + max(1, len(headers)))}{ws.max_row}"
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(45, max(12, max(len(str(c.value or "")) for c in column) + 2))
    output = BytesIO()
    wb.save(output)
    from django.http import HttpResponse
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response


def pdf_response(title, headers, rows, filename, iglesia, usuario):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_path = os.path.join(os.path.dirname(__import__("reportlab").__file__), "fonts", "Vera.ttf")
    if os.path.exists(font_path) and _FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(_FONT, font_path))
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    if _FONT in pdfmetrics.getRegisteredFontNames():
        styles["Title"].fontName = _FONT
        styles["Heading2"].fontName = _FONT
    body = [Paragraph(iglesia, styles["Title"]), Paragraph(title, styles["Heading2"]), Spacer(1, 8)]
    data = [headers] + [[str(v if v is not None else "") for v in row] for row in rows]
    if len(data) == 1:
        data.append(["Sin registros"] + [""] * (len(headers) - 1))
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1D4E89")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .25, colors.grey), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#EAF2F8")])]))
    body.append(table)
    doc.build(body)
    from django.http import HttpResponse
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
    return response
